"""Golden-file/snapshot test for the xlsx stub input pipeline.

Unlike test_xlsx_parser.py's unit tests (which assert on individual parsed
objects), this pins the FULL pipeline's output — xlsx bytes in,
generator/wiremock.py's real WireMock mapping JSON out — against a committed
fixture (tests/fixtures/xlsx_golden_mappings.json). A change to xlsx_parser.py
OR generator/wiremock.py that silently alters generated output for this
format fails here even if every individual unit test still passes in
isolation, because this is the one place both layers are exercised together
and compared against a fixed, human-reviewable expectation.

This fixture intentionally exercises several features in one pass:
  - a static (non-data-driven) single-scenario stub
  - a data-driven stub with a real JSONPath match condition + Default catch-all
  - a {brand}-templated URL (SOAP-shaped) — proving generator/wiremock.py's
    existing _apply_url_matcher already compiles "{word}" segments into a
    urlPattern regex with zero xlsx-specific handling needed (see the
    design-correction note in the xlsx implementation plan (internal notes, not in this repo) §4.1)
  - a %%Token%% response body, proving the ca_lisa_parser._resolve_variables
    reuse resolves it to Handlebars AND that has_dynamic_placeholders()
    correctly triggers the "response-template" transformer as a result

Regenerating the fixture after a deliberate behavior change: rerun the
generation block in this file's __main__ guard, review the diff by eye
(these are JSON mapping files — read them), then commit.
"""
from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl

from parser_worker.generator.wiremock import build_wiremock_mappings
from parser_worker.parsers import xlsx_parser as xp

_FIXTURE_PATH = Path(__file__).parent / "fixtures" / "xlsx_golden_mappings.json"

_STUBS_HEADER = [
    "Stub Name", "Model", "Protocol", "Method", "URL Path", "Templated URL?",
    "Request Body", "Response Status", "Response Content-Type",
    "Response Body  (inline, file:<name>, or SEE Rules)", "Data-driven?", "Notes",
]
_RULES_HEADER = [
    "Stub Name", "Order", "Scenario", "Extract Field", "Extract From",
    "Path / Expression", "Lookup File", "Match On", "Response Status",
    "Response Body (file:<name>)",
]


def _build_golden_workbook() -> bytes:
    wb = openpyxl.Workbook()
    stubs_ws = wb.active
    stubs_ws.title = "Stubs"
    stubs_ws.append(["STUBS"])
    stubs_ws.append(_STUBS_HEADER)
    stubs_ws.append(["SimpleStub", "Demo", "REST", "GET", "/v1/simple", "No", None, "200",
                      "application/json", '{"ok":true}', "No", ""])
    stubs_ws.append(["ComplexStub", "Demo", "REST", "POST", "/v1/complex", "No", None, "200",
                      "application/json", "SEE Rules tab", "Yes", ""])
    stubs_ws.append(["TemplatedStub", "Demo", "SOAP", "POST", "/{brand}/ACCTSVC120/01", "Yes", None, "200",
                      "text/xml", "<r><id>%%X-Interaction-Id%%</id></r>", "No", ""])

    rules_ws = wb.create_sheet("Rules")
    rules_ws.append(["RULES"])
    rules_ws.append(_RULES_HEADER)
    rules_ws.append(["ComplexStub", 1, "Success 1", "action", "body-json-path", "$.action",
                      None, "action == 'create'", 200, '{"status":"created"}'])
    rules_ws.append(["ComplexStub", 2, "Default", "action", "body-json-path", "$.action",
                      None, "always", 200, '{"status":"unknown"}'])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_mappings() -> list[dict]:
    xlsx_bytes = _build_golden_workbook()
    parsed, skip_notes = xp.parse_xlsx_zip(xlsx_bytes, {}, "golden-fixture.zip")
    assert skip_notes == [], f"golden fixture must parse cleanly, got skip notes: {skip_notes}"
    return [mapping for (_stub, _scenario, mapping) in build_wiremock_mappings(parsed)]


def test_generated_mappings_match_committed_golden_file():
    actual = _generate_mappings()
    expected = json.loads(_FIXTURE_PATH.read_text(encoding="utf-8"))

    assert len(actual) == expected["mapping_count"], (
        f"mapping count drifted: got {len(actual)}, fixture expects {expected['mapping_count']} — "
        "either a real behavior change (update the fixture deliberately, see this file's module "
        "docstring) or a regression."
    )
    assert actual == expected["mappings"], (
        "generated WireMock mappings no longer match tests/fixtures/xlsx_golden_mappings.json. "
        "If this is an intentional behavior change, regenerate the fixture and review the diff "
        "by eye before committing — don't just overwrite it blindly."
    )


if __name__ == "__main__":
    # Regenerate the committed fixture after a deliberate, reviewed change.
    mappings = _generate_mappings()
    _FIXTURE_PATH.write_text(
        json.dumps({"mapping_count": len(mappings), "mappings": mappings}, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(f"Regenerated {_FIXTURE_PATH} with {len(mappings)} mappings.")
