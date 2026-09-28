"""Tests for the Mockingbird XLSX stub template parser (Phase 0 — validator).

Fixtures are built in-memory with openpyxl rather than depending on the real
pilot sample file (still being filled in by the SV team as of this writing) —
per the implementation plan's stated default: build small, deterministic
fixtures now, swap in a golden-file test against the real sheet once it's
content-complete.
"""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest

from parser_worker.detector import detect_and_parse
from parser_worker.models import MatchType
from parser_worker.parsers import xlsx_parser as xp

STUBS_HEADER = [
    "Stub Name", "Model", "Protocol", "Method", "URL Path", "Templated URL?",
    "Request Body", "Response Status", "Response Content-Type",
    "Response Body  (inline, file:<name>, or SEE Rules)", "Data-driven?", "Notes",
]
RULES_HEADER = [
    "Stub Name", "Order", "Scenario", "Extract Field", "Extract From",
    "Path / Expression", "Lookup File", "Match On", "Response Status",
    "Response Body (file:<name>)",
]


def _build_workbook(stub_rows: list[list], rule_rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    stubs_ws = wb.active
    stubs_ws.title = "Stubs"
    stubs_ws.append(["STUBS"])  # row 1 — merged title row in the real template
    stubs_ws.append(STUBS_HEADER)  # row 2 — header
    for row in stub_rows:
        stubs_ws.append(row)

    rules_ws = wb.create_sheet("Rules")
    rules_ws.append(["RULES"])
    rules_ws.append(RULES_HEADER)
    for row in rule_rows:
        rules_ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in entries.items():
            zf.writestr(name, content)
    return buf.getvalue()


# A minimal, fully-clean template: one static stub, one data-driven stub with
# a real (non-placeholder) match condition and an explicit Default row.
def _clean_template_bytes() -> bytes:
    stubs = [
        ["SimpleStub", "Demo", "REST", "GET", "/v1/simple", "No", None, "200",
         "application/json", '{"ok":true}', "No", "1 scenario"],
        ["ComplexStub", "Demo", "REST", "POST", "/v1/complex", "No", None, "200",
         "application/json", "SEE Rules tab", "Yes", "2 scenarios"],
    ]
    rules = [
        ["ComplexStub", 1, "Success 1", "action", "body-json-path", "$.action",
         None, "action == 'create'", 200, "file:complex_s1_response.json"],
        ["ComplexStub", 2, "Default", "action", "body-json-path", "$.action",
         None, "always", 200, "file:complex_s2_response.json"],
    ]
    return _build_workbook(stubs, rules)


CLEAN_DATA_FILES = {
    "complex_s1_response.json": b'{"status":"created"}',
    "complex_s2_response.json": b'{"status":"unknown"}',
}


# ── sniff_xlsx_zip_entry ──────────────────────────────────────────────────────

class TestSniff:
    def test_detects_valid_template(self):
        z = _zip_bytes({"mockingbird-stub-template.xlsx": _clean_template_bytes()})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            assert xp.sniff_xlsx_zip_entry(zf) == "mockingbird-stub-template.xlsx"

    def test_ignores_xlsx_without_required_sheets(self):
        wb = openpyxl.Workbook()
        wb.active.title = "SomethingElse"
        buf = io.BytesIO()
        wb.save(buf)
        z = _zip_bytes({"not-a-stub-template.xlsx": buf.getvalue()})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            assert xp.sniff_xlsx_zip_entry(zf) is None

    def test_ignores_zip_with_no_xlsx(self):
        z = _zip_bytes({"a_Request_1.txt": b"={Method=\"GET\" URL=\"/x\"}"})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            assert xp.sniff_xlsx_zip_entry(zf) is None

    def test_ignores_zip_with_two_xlsx_files(self):
        template = _clean_template_bytes()
        z = _zip_bytes({"a.xlsx": template, "b.xlsx": template})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            assert xp.sniff_xlsx_zip_entry(zf) is None

    def test_malformed_xlsx_does_not_raise(self):
        z = _zip_bytes({"broken.xlsx": b"not a real xlsx file"})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            assert xp.sniff_xlsx_zip_entry(zf) is None


# ── check_zip_safety ──────────────────────────────────────────────────────────

class TestZipSafety:
    def test_clean_zip_passes(self):
        z = _zip_bytes({"a.xlsx": b"x", "data/b.json": b"{}"})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            assert xp.check_zip_safety(zf) is None

    def test_rejects_xlsm(self):
        z = _zip_bytes({"macros.xlsm": b"x"})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            reason = xp.check_zip_safety(zf)
            assert reason is not None and "macro" in reason.lower()

    def test_rejects_path_traversal_entry(self):
        z = _zip_bytes({"../../etc/passwd": b"x"})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            reason = xp.check_zip_safety(zf)
            assert reason is not None and "unsafe path" in reason.lower()

    def test_rejects_too_many_entries(self, monkeypatch):
        monkeypatch.setattr(xp, "MAX_ZIP_ENTRIES", 3)
        z = _zip_bytes({f"f{i}.txt": b"x" for i in range(5)})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            reason = xp.check_zip_safety(zf)
            assert reason is not None and "entries" in reason.lower()

    def test_rejects_oversized_uncompressed_total(self, monkeypatch):
        monkeypatch.setattr(xp, "MAX_ZIP_UNCOMPRESSED_BYTES", 10)
        z = _zip_bytes({"big.txt": b"x" * 1000})
        with zipfile.ZipFile(io.BytesIO(z)) as zf:
            reason = xp.check_zip_safety(zf)
            assert reason is not None and "limit" in reason.lower()


# ── validate_xlsx_zip ─────────────────────────────────────────────────────────

class TestValidate:
    def test_clean_template_is_valid_with_no_warnings(self):
        result = xp.validate_xlsx_zip(_clean_template_bytes(), CLEAN_DATA_FILES)
        assert result.valid is True
        assert result.errors == []
        assert result.warnings == []
        assert "2 stub(s)" in result.summary

    def test_missing_referenced_file_is_an_error(self):
        data_files = {"complex_s1_response.json": b"{}"}  # s2 missing
        result = xp.validate_xlsx_zip(_clean_template_bytes(), data_files)
        assert result.valid is False
        assert any("complex_s2_response.json" in str(e) for e in result.errors)

    def test_placeholder_path_expression_is_a_warning(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        rules = [["S1", 1, "Success 1", None, "body-json-path",
                   "(edit: request field/xpath that decides this branch)", None,
                   "request matches this case (edit condition)", 200, '{"x":1}']]
        result = xp.validate_xlsx_zip(_build_workbook(stubs, rules), {})
        assert result.valid is True  # no missing files -> not a hard error
        assert any("placeholder" in w.lower() for w in result.warnings)

    def test_orphaned_rules_row_is_a_warning(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        rules = [["TypoStubName", 1, "Default", None, "body-json-path", "$.x",
                   None, "always", 200, '{"x":1}']]
        result = xp.validate_xlsx_zip(_build_workbook(stubs, rules), {})
        assert any("does not appear on the Stubs tab" in w for w in result.warnings)
        # S1 itself is data-driven with zero matching rules -> also warned
        assert any("no matching rows on the Rules tab" in w for w in result.warnings)

    def test_data_driven_yes_with_no_rules_rows_is_a_warning(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        result = xp.validate_xlsx_zip(_build_workbook(stubs, []), {})
        assert any("no matching rows on the Rules tab" in w for w in result.warnings)

    def test_rules_rows_for_non_data_driven_stub_is_a_warning(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", '{"ok":true}', "No", ""]]
        rules = [["S1", 1, "Default", None, "body-json-path", "$.x", None,
                   "always", 200, '{"x":1}']]
        result = xp.validate_xlsx_zip(_build_workbook(stubs, rules), {})
        assert any("Data-driven? = No" in w for w in result.warnings)

    def test_missing_catch_all_is_a_warning(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        rules = [["S1", 1, "Success 1", None, "body-json-path", "$.x", None,
                   "x == 'y'", 200, '{"x":1}']]
        result = xp.validate_xlsx_zip(_build_workbook(stubs, rules), {})
        assert any("no Default/catch-all" in w for w in result.warnings)

    def test_legacy_token_in_content_type_is_a_warning(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "%%Content-Type%%", '{"ok":true}', "No", ""]]
        result = xp.validate_xlsx_zip(_build_workbook(stubs, []), {})
        assert any("%%Content-Type%%" in w for w in result.warnings)

    def test_numeric_looking_csv_header_is_a_warning(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        rules = [["S1", 1, "Success 1", "action", "body-json-path", "$.action",
                   "lookup.csv", "lookup: action == success", 200, '{"x":1}'],
                  ["S1", 2, "Default", "action", "body-json-path", "$.action",
                   "lookup.csv", "always", 200, '{"x":2}']]
        data_files = {"lookup.csv": b"833,action,flow#\nCLRF1,UPDT,success1\n"}
        result = xp.validate_xlsx_zip(_build_workbook(stubs, rules), data_files)
        assert any("looks like data, not a column name" in w for w in result.warnings)

    def test_missing_required_sheet_is_a_hard_error(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Stubs"
        wb.active.append(STUBS_HEADER)
        buf = io.BytesIO()
        wb.save(buf)
        result = xp.validate_xlsx_zip(buf.getvalue(), {})
        assert result.valid is False
        assert any("Rules" in str(e) for e in result.errors)

    def test_unreadable_workbook_is_a_hard_error(self):
        result = xp.validate_xlsx_zip(b"not an xlsx file at all", {})
        assert result.valid is False
        assert len(result.errors) == 1


# ── parse_xlsx_zip (Phase 1) ──────────────────────────────────────────────────

class TestParse:
    def test_clean_template_builds_both_stubs(self):
        parsed, skip_notes = xp.parse_xlsx_zip(_clean_template_bytes(), CLEAN_DATA_FILES, "test.zip")
        assert {s.name for s in parsed.stubs} == {"SimpleStub", "ComplexStub"}
        assert skip_notes == []

    def test_lookup_file_stub_is_skipped_when_csv_missing_not_misrepresented(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        rules = [["S1", 1, "Success 1", "bban", "body-xpath", "(edit: request field/xpath that decides this branch)",
                   "lookup.csv", "lookup: bban == success", 200, '{"x":1}'],
                  ["S1", 2, "Default", "bban", "body-xpath", "(edit: request field/xpath that decides this branch)",
                   "lookup.csv", "always", 200, '{"x":2}']]
        wb = _build_workbook(stubs, rules)
        parsed, skip_notes = xp.parse_xlsx_zip(wb, {}, "test.zip")
        assert parsed.stubs == []
        assert any("lookup.csv" in n and "not in the uploaded zip" in n for n in skip_notes)

    def test_lookup_file_stub_with_single_key_routes_correctly(self):
        # Mirrors the real pilot LKP01_RetrieveAccountStatus shape exactly:
        # one key column, a Default catch-all, and CSV outcome labels
        # ("success"/"error") that only loosely match the numbered Rules
        # scenario names ("Success 1"/"Error 2") — exercising the
        # trailing-digit-stripped resolution path, not just exact match.
        stubs = [["LKP01_RetrieveAccountStatus", "ESP Model", "SOAP", "POST", "/{brand}/ACCTSVC120/01",
                   "Yes", None, "200", "text/xml", "SEE Rules tab", "Yes", ""]]
        rules = [
            ["LKP01_RetrieveAccountStatus", 1, "Success 1", "bban", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP01.csv",
             "lookup: bban == success", 200, "file:s1.xml"],
            ["LKP01_RetrieveAccountStatus", 2, "Error 2", "bban", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP01.csv",
             "lookup: bban == error", 200, "file:s2.xml"],
            ["LKP01_RetrieveAccountStatus", 3, "Default", "bban", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP01.csv",
             "always", 200, "file:s3.xml"],
        ]
        data_files = {
            "LKP01.csv": b"bban,scenario\n90000000005357,error\n90000000005356,success\n",
            "s1.xml": b"<r>success-body</r>",
            "s2.xml": b"<r>error-body</r>",
            "s3.xml": b"<r>default-body</r>",
        }
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, rules), data_files, "test.zip")
        assert skip_notes == []
        assert len(parsed.stubs) == 1
        stub = parsed.stubs[0]
        assert stub.force_lookup_table is True
        assert stub.lookup_discriminator_type == "xpath"
        assert stub.lookup_discriminator_field == "bban"

        by_key = {s.lookup_key: s for s in stub.scenarios if s.lookup_key is not None}
        assert by_key["90000000005356"].body == "<r>success-body</r>"
        assert by_key["90000000005357"].body == "<r>error-body</r>"
        unkeyed = [s for s in stub.scenarios if s.lookup_key is None]
        assert len(unkeyed) == 1
        assert unkeyed[0].body == "<r>default-body</r>"
        assert unkeyed[0].match.type == MatchType.ALWAYS

    def test_lookup_file_stub_with_composite_key_routes_correctly(self):
        # The corrected version of the LKP03 gap above: 'Extract Field'
        # comma-separates BOTH key columns, matching LKP03.csv's real
        # 2-column-key shape exactly (once the sheet is fixed to name both).
        stubs = [["LKP03_InitiateUpdate", "ESP Model", "SOAP", "POST", "/{brand}/ACCTRNPRC090/03",
                   "Yes", None, "200", "text/xml", "SEE Rules tab", "Yes", ""]]
        rules = [
            ["LKP03_InitiateUpdate", 1, "Success 1", "identifier,action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP03.csv",
             "lookup: action == success", 200, "file:s1.xml"],
            ["LKP03_InitiateUpdate", 2, "Success 2", "identifier,action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP03.csv",
             "lookup: action == success", 200, "file:s2.xml"],
            ["LKP03_InitiateUpdate", 3, "Error 3", "identifier,action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP03.csv",
             "lookup: action == error", 200, "file:s3.xml"],
            ["LKP03_InitiateUpdate", 4, "Default", "identifier,action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP03.csv",
             "always", 200, "file:s4.xml"],
        ]
        data_files = {
            "LKP03.csv": (
                b"identifier,action,flow\n"
                b"CLRF7458USDA23,UPDT,success1\n"
                b"CLRF7458USDA24,UPDT,success2\n"
                b"CLRF7458USDA25,UPDT,error\n"
            ),
            "s1.xml": b"<r>1</r>", "s2.xml": b"<r>2</r>", "s3.xml": b"<r>3</r>", "s4.xml": b"<r>default</r>",
        }
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, rules), data_files, "test.zip")
        assert skip_notes == []
        assert len(parsed.stubs) == 1
        stub = parsed.stubs[0]
        assert stub.lookup_discriminator_field == "identifier,action"

        by_key = {s.lookup_key: s for s in stub.scenarios if s.lookup_key is not None}
        assert len(by_key) == 3
        assert by_key["CLRF7458USDA23\x1fUPDT"].body == "<r>1</r>"
        assert by_key["CLRF7458USDA24\x1fUPDT"].body == "<r>2</r>"
        assert by_key["CLRF7458USDA25\x1fUPDT"].body == "<r>3</r>"
        # Composite match condition combines both fields with AND, for the
        # no-Java-extension static-mapping fallback path.
        assert "identifier" in by_key["CLRF7458USDA23\x1fUPDT"].match.value
        assert "action" in by_key["CLRF7458USDA23\x1fUPDT"].match.value
        assert " and " in by_key["CLRF7458USDA23\x1fUPDT"].match.value

    def test_lookup_file_stub_skipped_when_extract_field_misses_a_key_column(self):
        # Mirrors the real pilot LKP03_InitiateUpdate gap exactly: the
        # CSV has 2 key columns (identifier + action) but 'Extract Field'
        # only names one ("action") — genuinely insufficient to build a
        # correct composite key, not something to guess around.
        stubs = [["LKP03_InitiateUpdate", "ESP Model", "SOAP", "POST", "/{brand}/ACCTRNPRC090/03",
                   "Yes", None, "200", "text/xml", "SEE Rules tab", "Yes", ""]]
        rules = [
            ["LKP03_InitiateUpdate", 1, "Success 1", "action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP03.csv",
             "lookup: action == success", 200, "file:s1.xml"],
            ["LKP03_InitiateUpdate", 2, "Default", "action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "LKP03.csv",
             "always", 200, "file:s2.xml"],
        ]
        data_files = {
            "LKP03.csv": b"identifier,action,flow\nCLRF7458USDA23,UPDT,success1\nCLRF7458USDA24,UPDT,success2\n",
        }
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, rules), data_files, "test.zip")
        assert parsed.stubs == []
        assert any("2 key column" in n and "1 field name" in n for n in skip_notes)

    def test_lookup_file_ambiguous_outcome_label_is_skipped_not_guessed(self):
        # A CSV outcome value ("success") that matches more than one
        # similarly-named Rules scenario ("Success 1" and "Success 2")
        # equally well must never be resolved by a coin flip.
        stubs = [["S1", "Demo", "SOAP", "POST", "/v1/s1", "No", None, "200",
                   "text/xml", "SEE Rules tab", "Yes", ""]]
        rules = [
            ["S1", 1, "Success 1", "action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "lookup.csv",
             "lookup: action == success", 200, "file:s1.xml"],
            ["S1", 2, "Success 2", "action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "lookup.csv",
             "lookup: action == success", 200, "file:s2.xml"],
            ["S1", 3, "Default", "action", "body-xpath",
             "(edit: request field/xpath that decides this branch)", "lookup.csv",
             "always", 200, "file:s3.xml"],
        ]
        data_files = {
            "lookup.csv": b"action,outcome\nCANC,success\n",
            "s1.xml": b"<r>1</r>", "s2.xml": b"<r>2</r>", "s3.xml": b"<r>default</r>",
        }
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, rules), data_files, "test.zip")
        assert len(parsed.stubs) == 1
        # Only the unresolvable CSV row is skipped — the Default catch-all
        # still generates, so the stub as a whole isn't silently dropped.
        assert len(parsed.stubs[0].scenarios) == 1
        assert parsed.stubs[0].scenarios[0].match.type == MatchType.ALWAYS
        assert any("did not resolve to exactly one Rules scenario" in n for n in skip_notes)

    def test_placeholder_only_stub_is_skipped(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        rules = [["S1", 1, "Success 1", None, "body-json-path",
                   "(edit: request field/xpath that decides this branch)", None,
                   "request matches this case (edit condition)", 200, '{"x":1}']]
        wb = _build_workbook(stubs, rules)
        parsed, skip_notes = xp.parse_xlsx_zip(wb, {}, "test.zip")
        assert parsed.stubs == []
        assert any("no usable scenario rows" in n for n in skip_notes)

    def test_unsupported_protocol_is_skipped(self):
        stubs = [["S1", "Demo", "Kafka", "POST", "/v1/s1", "No", None, "200",
                   "application/json", '{"ok":true}', "No", ""]]
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, []), {}, "test.zip")
        assert parsed.stubs == []
        assert any("Kafka/MQ JSON formats" in n for n in skip_notes)

    def test_unrecognised_method_is_skipped(self):
        stubs = [["S1", "Demo", "REST", "FETCH", "/v1/s1", "No", None, "200",
                   "application/json", '{"ok":true}', "No", ""]]
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, []), {}, "test.zip")
        assert parsed.stubs == []
        assert any("unrecognised HTTP method" in n for n in skip_notes)

    def test_missing_response_file_is_skipped_not_crashed(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "file:missing.json", "No", ""]]
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, []), {}, "test.zip")
        assert parsed.stubs == []
        assert any("missing.json" in n for n in skip_notes)

    def test_legacy_token_in_body_resolved_to_handlebars(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", '{"id":"%%X-Interaction-Id%%"}', "No", ""]]
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, []), {}, "test.zip")
        assert skip_notes == []
        assert parsed.stubs[0].scenarios[0].body == '{"id":"{{request.headers.X-Interaction-Id}}"}'

    def test_scenario_order_follows_the_order_column(self):
        stubs = [["S1", "Demo", "REST", "GET", "/v1/s1", "No", None, "200",
                   "application/json", "SEE Rules tab", "Yes", ""]]
        # Deliberately out of order in the sheet — Order column must win.
        rules = [["S1", 2, "Second", None, "body-json-path", "$.x", None, "x == 'b'", 200, '{"x":"b"}'],
                  ["S1", 1, "First", None, "body-json-path", "$.x", None, "x == 'a'", 200, '{"x":"a"}'],
                  ["S1", 3, "Default", None, "body-json-path", "$.x", None, "always", 200, '{"x":"z"}']]
        parsed, skip_notes = xp.parse_xlsx_zip(_build_workbook(stubs, rules), {}, "test.zip")
        assert skip_notes == []
        assert [s.name for s in parsed.stubs[0].scenarios] == ["First", "Second", "Default"]


# ── end-to-end through detector.detect_and_parse ──────────────────────────────

class TestDetectorIntegration:
    def test_xlsx_zip_routes_to_xlsx_and_builds_real_stubs(self, tmp_path):
        z = _zip_bytes({
            "mockingbird-stub-template.xlsx": _clean_template_bytes(),
            "data/complex_s1_response.json": CLEAN_DATA_FILES["complex_s1_response.json"],
            "data/complex_s2_response.json": CLEAN_DATA_FILES["complex_s2_response.json"],
        })
        zip_path = tmp_path / "upload.zip"
        zip_path.write_bytes(z)

        parser, result, parsed_file = detect_and_parse(zip_path)

        assert result.format_detected == "mockingbird-xlsx-stub-template"
        assert result.valid is True
        assert parsed_file is not None
        assert {s.name for s in parsed_file.stubs} == {"SimpleStub", "ComplexStub"}

        simple = next(s for s in parsed_file.stubs if s.name == "SimpleStub")
        assert len(simple.scenarios) == 1
        assert simple.scenarios[0].match.type.value == "always"
        assert simple.scenarios[0].body == '{"ok":true}'

        complex_stub = next(s for s in parsed_file.stubs if s.name == "ComplexStub")
        assert [sc.name for sc in complex_stub.scenarios] == ["Success 1", "Default"]
        assert complex_stub.scenarios[0].match.type.value == "body-json-path"
        assert complex_stub.scenarios[0].match.value == "$.action"
        assert complex_stub.scenarios[1].match.type.value == "always"
        assert complex_stub.scenarios[0].response_headers == {"Content-Type": "application/json"}

    def test_data_files_resolve_regardless_of_zip_nesting(self, tmp_path):
        # Files at the zip root (not under data/) must still resolve by basename.
        z = _zip_bytes({
            "template.xlsx": _clean_template_bytes(),
            "complex_s1_response.json": CLEAN_DATA_FILES["complex_s1_response.json"],
            "complex_s2_response.json": CLEAN_DATA_FILES["complex_s2_response.json"],
        })
        zip_path = tmp_path / "upload.zip"
        zip_path.write_bytes(z)

        _, result, _ = detect_and_parse(zip_path)
        assert result.valid is True
        assert result.errors == []

    def test_ca_lisa_zip_still_routes_unchanged(self, tmp_path):
        z = _zip_bytes({
            "Op_Request_20260610_100059.txt": b'={Method="GET" URL="/x" httpDetails={Version="1.1" httpHeaders={}}}',
            "Op_Response_20260610_100059.txt": b'ResponseHeader={StatusCode="200"}\nResponse..{"ok":true}',
        })
        zip_path = tmp_path / "upload.zip"
        zip_path.write_bytes(z)

        parser, result, parsed_file = detect_and_parse(zip_path)

        assert result.format_detected == "ca-lisa-http-pair"
        assert result.valid is True
        assert parsed_file is not None
        assert len(parsed_file.stubs) == 1

    def test_macro_enabled_workbook_is_rejected_with_a_clear_message(self, tmp_path):
        z = _zip_bytes({"template.xlsm": _clean_template_bytes()})
        zip_path = tmp_path / "upload.zip"
        zip_path.write_bytes(z)

        _, result, parsed_file = detect_and_parse(zip_path)
        assert result.valid is False
        assert parsed_file is None
        assert any("macro" in str(e).lower() for e in result.errors)
